import PyPDF2
from docx import Document as DocxDocument
from pathlib import Path
from typing import List
from .types import DocumentChunk, ParsedDocument, DocumentType

class DocumentParser:
    CHUNK_SIZE = 2000  # Characters per chunk
    CHUNK_OVERLAP = 200

    @staticmethod
    def parse_file(file_path: str) -> ParsedDocument:
        """Parse a file (PDF or DOCX) and return structured chunks"""
        path = Path(file_path)
        
        if path.suffix.lower() == '.pdf':
            return DocumentParser._parse_pdf(file_path)
        elif path.suffix.lower() in ['.docx', '.doc']:
            return DocumentParser._parse_docx(file_path)
        else:
            raise ValueError(f"Unsupported file type: {path.suffix}")

    @staticmethod
    def _parse_pdf(file_path: str) -> ParsedDocument:
        """Extract text from PDF and chunk it"""
        chunks: List[DocumentChunk] = []
        
        with open(file_path, 'rb') as file:
            pdf_reader = PyPDF2.PdfReader(file)
            total_pages = len(pdf_reader.pages)
            full_text = ""
            
            for page_num, page in enumerate(pdf_reader.pages):
                text = page.extract_text()
                full_text += f"\n--- Page {page_num + 1} ---\n{text}"

        # Chunk the text
        chunks = DocumentParser._chunk_text(full_text, Path(file_path).name)
        
        return ParsedDocument(
            original_filename=Path(file_path).name,
            document_type=DocumentType.PDF,
            total_pages=total_pages,
            chunks=chunks,
            metadata={'file_type': 'pdf', 'pages': total_pages}
        )

    @staticmethod
    def _parse_docx(file_path: str) -> ParsedDocument:
        """Extract text from DOCX and chunk it"""
        doc = DocxDocument(file_path)
        full_text = "\n".join([para.text for para in doc.paragraphs])
        
        chunks = DocumentParser._chunk_text(full_text, Path(file_path).name)
        
        return ParsedDocument(
            original_filename=Path(file_path).name,
            document_type=DocumentType.DOCX,
            total_pages=1,
            chunks=chunks,
            metadata={'file_type': 'docx', 'paragraphs': len(doc.paragraphs)}
        )

    @staticmethod
    def _chunk_text(text: str, source: str) -> List[DocumentChunk]:
        """Split text into overlapping chunks"""
        chunks: List[DocumentChunk] = []
        chunk_size = DocumentParser.CHUNK_SIZE
        overlap = DocumentParser.CHUNK_OVERLAP
        
        for i in range(0, len(text), chunk_size - overlap):
            chunk_content = text[i:i + chunk_size]
            if chunk_content.strip():
                chunk = DocumentChunk(
                    content=chunk_content,
                    source=source,
                    page=1,
                    chunk_index=len(chunks)
                )
                chunks.append(chunk)
        
        return chunks
